"""Excel import/export following the ActBasedDigTwinFrmwrk workbook layout:
one sheet per view (goals, seds, utility_types, kpis, plans, actuals,
org_units, employees, equipment, resources, event_kinds, erp_objects,
erp_mappings, initial_state)."""

from __future__ import annotations

import io
from typing import Any, BinaryIO

from openpyxl import Workbook, load_workbook

from dtb.domain.model import EnterpriseModel
from dtb.integration.base import ImportReport
from dtb.integration.tabular import SHEETS, apply_rows, model_to_rows

# Sheet name aliases (Russian names from the source workbook are accepted).
_ALIASES = {
    "цели": "goals",
    "сэд": "seds",
    "сэды": "seds",
    "полезности": "utility_types",
    "ресурсы": "resources",
    "оргструктура": "org_units",
    "сотрудники": "employees",
    "оборудование": "equipment",
    "события": "event_kinds",
    "kpi": "kpis",
    "показатели": "kpis",
    "erp": "erp_objects",
    "связи": "erp_mappings",
    "факт": "actuals",
    "план": "plans",
    "планы": "plans",
    "прогноз": "plans",
    "состояние": "initial_state",
}


def import_workbook(model: EnterpriseModel, data: bytes | BinaryIO) -> ImportReport:
    report = ImportReport(source="excel")
    stream = io.BytesIO(data) if isinstance(data, bytes) else data
    try:
        wb = load_workbook(stream, data_only=True, read_only=True)
    except Exception as exc:  # noqa: BLE001
        report.error("workbook", f"не удалось открыть файл: {exc}")
        return report
    # apply in dependency-friendly order regardless of workbook sheet order
    found: dict[str, list[dict[str, Any]]] = {}
    for ws in wb.worksheets:
        sheet = _ALIASES.get(ws.title.strip().lower(), ws.title.strip().lower())
        if sheet not in SHEETS:
            report.warn(ws.title, "лист не распознан — пропущен")
            continue
        rows_iter = ws.iter_rows(values_only=True)
        try:
            header = [str(h or "").strip().lower() for h in next(rows_iter)]
        except StopIteration:
            continue
        rows = [dict(zip(header, row)) for row in rows_iter]
        found[sheet] = rows
    for sheet in SHEETS:
        if sheet in found:
            apply_rows(model, sheet, found[sheet], report)
    wb.close()
    return report


def export_workbook(model: EnterpriseModel) -> bytes:
    """Export the model as a workbook (also serves as the import template)."""
    wb = Workbook()
    wb.remove(wb.active)
    for sheet in SHEETS:
        ws = wb.create_sheet(title=sheet)
        header, rows = model_to_rows(model, sheet)
        ws.append(header)
        for row in rows:
            ws.append(row)
        ws.freeze_panes = "A2"
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()
