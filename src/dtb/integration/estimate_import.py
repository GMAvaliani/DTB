"""Импорт сметы со статусом выполнения в формате «объекты → разделы → позиции».

Поддерживается распространённая форма строительной сметы:

* строка заголовков со столбцами «Наименование работ», «Обозн.», «Объем»,
  «Материал ед.», «Материал всего», «Работа ед.», «Работа всего», «Доставка»,
  «Аренда техники», «Примечание»;
* справа — блок «Выполнение на ДД.ММ.ГГГГ» с теми же столбцами факта:
  «Объем», «Материал всего», «Работа всего»;
* объекты (здания) и разделы (нулевой цикл, КМ, кровля…) — строки без единиц
  измерения; итоговые строки помечены «ИТОГО» / «Всего»;
* внизу — сведения об авансе: «аванс» и «зачет аванса NN%».

Распознавание эвристическое, поэтому импорт возвращает подробный отчёт: что
стало объектом, что разделом, какие строки пропущены и где арифметика сметы
не сходится (объём × цена ≠ итог). Итоги сметы считаются бюджетом как есть —
расхождения не исправляются молча, а показываются.
"""

from __future__ import annotations

import io
import re
from dataclasses import dataclass, field
from datetime import date, datetime
from typing import Any

from openpyxl import load_workbook

from dtb.domain.model import EnterpriseModel
from dtb.domain.projects import Project, WorkPackage
from dtb.integration.base import ImportReport

#: Подписи итоговых строк.
_TOTAL_MARKERS = ("итого", "всего")
#: Единицы, при которых работа считается «по времени» (level of effort).
_LOE_UNITS = ("мес", "мес.", "месяц", "дн", "дн.", "день", "сут", "нед", "нед.")
#: Объект по умолчанию для позиций до первого объекта (подготовка площадки).
DEFAULT_OBJECT = "Подготовительные работы"


@dataclass
class EstimateImportResult:
    project: Project
    work_packages: list[WorkPackage]
    report: ImportReport
    summary: dict[str, Any] = field(default_factory=dict)

    def to_dict(self) -> dict[str, Any]:
        return {
            "project": self.project.model_dump(mode="json"),
            "work_packages": len(self.work_packages),
            "report": self.report.to_dict(),
            "summary": self.summary,
        }


def _text(value: Any) -> str:
    return "" if value is None else re.sub(r"\s+", " ", str(value)).strip()


def _num(value: Any) -> float | None:
    if value is None or isinstance(value, bool):
        return None
    if isinstance(value, (int, float)):
        return float(value)
    text = _text(value).replace(" ", "").replace(",", ".")
    try:
        return float(text)
    except ValueError:
        return None


def _find_header(rows: list[tuple]) -> int | None:
    for index, row in enumerate(rows[:30]):
        if any(_text(c).lower() == "наименование работ" for c in row):
            return index
    return None


def _column_map(header: tuple) -> dict[str, int]:
    """Номера столбцов плана и факта по подписям заголовка."""
    names = [_text(c).lower() for c in header]
    cols: dict[str, int] = {}

    def first(label: str, start: int = 0) -> int | None:
        for i in range(start, len(names)):
            if names[i] == label:
                return i
        return None

    mapping = {
        "name": "наименование работ",
        "unit": "обозн.",
        "qty": "объем",
        "mat_unit": "материал ед.",
        "mat_total": "материал всего",
        "lab_unit": "работа ед.",
        "lab_total": "работа всего",
        "delivery": "доставка",
        "equipment": "аренда техники",
        "note": "примечание",
    }
    for key, label in mapping.items():
        found = first(label)
        if found is not None:
            cols[key] = found
    # блок факта начинается правее «Примечания» либо правее «Аренды техники»
    status_from = max(cols.get("note", -1), cols.get("equipment", -1), cols.get("lab_total", -1)) + 1
    for key, label in (("act_qty", "объем"), ("act_mat", "материал всего"),
                       ("act_lab", "работа всего"), ("remaining", "остаток (объем)")):
        found = first(label, status_from)
        if found is not None:
            cols[key] = found
    return cols


def _status_date(rows: list[tuple], header_index: int) -> tuple[str, date | None]:
    pattern = re.compile(r"(\d{1,2})\.(\d{1,2})\.(\d{4})")
    for row in rows[: header_index + 1]:
        for cell in row:
            text = _text(cell)
            if "выполнение" in text.lower():
                match = pattern.search(text)
                if match:
                    d, m, y = (int(x) for x in match.groups())
                    try:
                        return text, date(y, m, d)
                    except ValueError:
                        return text, None
                return text, None
    return "", None


def _cell(row: tuple, cols: dict[str, int], key: str) -> Any:
    index = cols.get(key)
    if index is None or index >= len(row):
        return None
    return row[index]


def _kind(row: tuple, cols: dict[str, int]) -> str:
    """Тип строки: item | header | total | blank."""
    name = _text(_cell(row, cols, "name"))
    unit = _text(_cell(row, cols, "unit"))
    qty_cell = _cell(row, cols, "qty")
    if _text(qty_cell).lower() in _TOTAL_MARKERS:
        return "total"
    if not name:
        return "blank"
    qty = _num(qty_cell)
    if unit and qty is not None and qty > 0:
        return "item"
    return "header"


def parse_estimate(
    data: bytes,
    project_id: str,
    project_name: str = "",
    sheet: str | None = None,
    start_date: date | None = None,
) -> EstimateImportResult:
    """Разобрать смету со статусом выполнения в проект и пакеты работ."""
    report = ImportReport(source="estimate")
    wb = load_workbook(io.BytesIO(data), data_only=True, read_only=True)
    worksheets = [wb[sheet]] if sheet else list(wb.worksheets)
    ws = None
    rows: list[tuple] = []
    header_index = None
    for candidate in worksheets:
        candidate_rows = [tuple(r) for r in candidate.iter_rows(values_only=True)]
        header_index = _find_header(candidate_rows)
        if header_index is not None:
            ws, rows = candidate, candidate_rows
            break
    wb.close()
    project = Project(id=project_id, name=project_name or project_id)
    if ws is None or header_index is None:
        report.error("estimate", "не найдена строка заголовков со столбцом «Наименование работ»")
        return EstimateImportResult(project, [], report)

    cols = _column_map(rows[header_index])
    for required in ("name", "unit", "qty"):
        if required not in cols:
            report.error("estimate", f"нет обязательного столбца: {required}")
    if report.errors:
        return EstimateImportResult(project, [], report)

    status_label, status_day = _status_date(rows, header_index)
    has_status = "act_qty" in cols
    body = rows[header_index + 1:]
    kinds = [_kind(r, cols) for r in body]

    def next_meaningful(i: int) -> str | None:
        for k in kinds[i + 1:]:
            if k in ("item", "header"):
                return k
        return None

    packages: list[WorkPackage] = []
    current_object = DEFAULT_OBJECT
    current_section = "Общие"
    seen_object = False
    objects_order: list[str] = [DEFAULT_OBJECT]
    object_counter: dict[str, int] = {DEFAULT_OBJECT: 0}
    advance = 0.0
    offset_share = 0.0
    excel_row0 = header_index + 2  # номер строки Excel для первой строки тела

    for i, (row, kind) in enumerate(zip(body, kinds)):
        excel_row = excel_row0 + i
        # сведения об авансе встречаются в любых столбцах
        for j, cell in enumerate(row):
            label = _text(cell).lower()
            if label == "аванс" and j + 1 < len(row) and _num(row[j + 1]) is not None:
                advance = float(_num(row[j + 1]) or 0.0)
            match = re.match(r"зач[её]т аванса\s*(\d+(?:[.,]\d+)?)\s*%", label)
            if match:
                offset_share = float(match.group(1).replace(",", ".")) / 100.0
        if kind == "header":
            name = _text(_cell(row, cols, "name"))
            following = next_meaningful(i)
            if following == "header":
                current_object = name
                current_section = "Общие"
                seen_object = True
                if name not in object_counter:
                    objects_order.append(name)
                    object_counter[name] = 0
                report.warn("estimate", f"объект: {name}", row=excel_row)
            elif not seen_object or "не посчитано" in name.lower():
                report.warn(
                    "estimate",
                    f"строка без объёма и цены пропущена: {name}",
                    row=excel_row,
                )
            else:
                current_section = name
            continue
        if kind != "item":
            continue

        name = _text(_cell(row, cols, "name"))
        unit = _text(_cell(row, cols, "unit"))
        qty = float(_num(_cell(row, cols, "qty")) or 0.0)
        mat_unit = _num(_cell(row, cols, "mat_unit")) or 0.0
        lab_unit = _num(_cell(row, cols, "lab_unit")) or 0.0
        mat_total = _num(_cell(row, cols, "mat_total"))
        lab_total = _num(_cell(row, cols, "lab_total"))
        for label, unit_cost, total in (
            ("материалы", mat_unit, mat_total),
            ("работы", lab_unit, lab_total),
        ):
            if unit_cost and total is not None:
                expected = qty * unit_cost
                if abs(expected - total) > max(1.0, 0.01 * abs(expected)):
                    report.warn(
                        "estimate",
                        f"{name[:60]}: {label} — объём × цена = {expected:,.2f}, "
                        f"а в смете {total:,.2f}; в бюджет взят итог сметы",
                        row=excel_row,
                    )
        object_counter[current_object] += 1
        obj_index = objects_order.index(current_object)
        wp_id = f"{project_id}_o{obj_index:02d}_{object_counter[current_object]:03d}"
        act_qty = act_mat = act_lab = 0.0
        if has_status:
            act_qty = max(0.0, float(_num(_cell(row, cols, "act_qty")) or 0.0))
            act_mat = max(0.0, float(_num(_cell(row, cols, "act_mat")) or 0.0))
            act_lab = max(0.0, float(_num(_cell(row, cols, "act_lab")) or 0.0))
            if act_qty > qty * 1.01 and qty > 0:
                report.warn(
                    "estimate",
                    f"{name[:60]}: выполнено {act_qty:g} {unit} больше сметных "
                    f"{qty:g} — перерасход объёма",
                    row=excel_row,
                )
        packages.append(
            WorkPackage(
                id=wp_id,
                name=name,
                project_ref=project_id,
                object_name=current_object,
                section=current_section,
                unit=unit,
                planned_quantity=qty,
                material_unit_cost=mat_unit,
                labor_unit_cost=lab_unit,
                material_total=mat_total if mat_total is not None else None,
                labor_total=lab_total if lab_total is not None else None,
                delivery_cost=float(_num(_cell(row, cols, "delivery")) or 0.0),
                equipment_cost=float(_num(_cell(row, cols, "equipment")) or 0.0),
                actual_quantity=act_qty,
                actual_material_cost=act_mat,
                actual_labor_cost=act_lab,
                level_of_effort=unit.lower() in _LOE_UNITS,
            )
        )
        report.count("work_package", True)

    project.advance_payment = advance
    project.advance_offset_share = offset_share
    project.status_label = status_label
    if status_day is not None and start_date is not None:
        project.status_date = max(0.0, (status_day - start_date).days * 24.0)

    by_object: dict[str, dict[str, float]] = {}
    for wp in packages:
        row = by_object.setdefault(
            wp.object_name, {"positions": 0, "bac": 0.0, "actual_cost": 0.0, "earned": 0.0}
        )
        row["positions"] += 1
        row["bac"] += wp.bac()
        row["actual_cost"] += wp.actual_cost()
        row["earned"] += wp.bac() * wp.actual_fraction()
    summary = {
        "sheet": ws.title,
        "status_label": status_label,
        "status_date": status_day.isoformat() if status_day else None,
        "objects": {
            k: {kk: round(vv, 2) for kk, vv in v.items()} for k, v in by_object.items()
        },
        "bac": round(sum(wp.bac() for wp in packages), 2),
        "actual_cost": round(sum(wp.actual_cost() for wp in packages), 2),
        "earned_value": round(sum(wp.bac() * wp.actual_fraction() for wp in packages), 2),
        "advance_payment": advance,
        "advance_offset_share": offset_share,
        "level_of_effort": sum(1 for wp in packages if wp.level_of_effort),
    }
    return EstimateImportResult(project, packages, report, summary)


def import_estimate(
    model: EnterpriseModel,
    data: bytes,
    project_id: str,
    project_name: str = "",
    sheet: str | None = None,
    chain_within_object: bool = True,
) -> EstimateImportResult:
    """Разобрать смету и добавить проект с пакетами работ в модель.

    ``chain_within_object`` — связать позиции одного объекта последовательно
    (каждая начинается после предыдущей). Работы «по времени» и подготовка
    площадки остаются без предшественников. Связи затем правятся вручную.
    """
    start = None
    if model.enterprise.start_date:
        try:
            start = datetime.fromisoformat(model.enterprise.start_date).date()
        except ValueError:
            start = None
    result = parse_estimate(data, project_id, project_name, sheet, start)
    if result.report.errors:
        return result
    if chain_within_object:
        previous: dict[str, str] = {}
        for wp in result.work_packages:
            if wp.level_of_effort or wp.object_name == DEFAULT_OBJECT:
                continue
            if wp.object_name in previous:
                wp.predecessors = [previous[wp.object_name]]
            previous[wp.object_name] = wp.id
    model.projects[result.project.id] = result.project
    for wp in result.work_packages:
        model.work_packages[wp.id] = wp
    return result
